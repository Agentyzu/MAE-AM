from collections import defaultdict
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import unicodedata

import numpy as np

from utils.text_length import count_text

POLICY = "category_ten_length_intervals_without_replacement"
INTERVALS = 10
_ATVI_ALIASES = {"alcohol & spirits": "Alcohol", "food & beverage": "Food",
                 "health care": "Health", "home & personal care": "Home",
                 "media & entertainment": "Media"}


def identity(text):
    return " ".join(unicodedata.normalize("NFKC", text).split()).casefold()


def _file_info(path, code):
    return {"path": path.relative_to(code).as_posix(),
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def load_pools(code, datasets):
    code = Path(code)
    pools, excluded = [], []
    for dataset in datasets:
        configured = sorted(path.name[len("config_"):-len("_Prom.py")]
                            for path in (code / f"cfgs/{dataset}_Prom_configs").glob("config_*_Prom.py"))
        rows, sources, labels = defaultdict(list), defaultdict(list), defaultdict(set)
        workbook = code / "dataset/ATVI/ad_data.xlsx"
        if dataset == "ATVI" and workbook.exists():
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise RuntimeError("Install requirements-experiments.txt to read original ATVI data") from exc
            book = load_workbook(workbook, read_only=True, data_only=True)
            try:
                sheet = book.worksheets[0]
                iterator = sheet.iter_rows(values_only=True)
                header = list(next(iterator))
                columns = {key: header.index(key) for key in ("Category", "Advertiser", "Ad_copy")}
                known = {identity(name): name for name in configured}
                omitted = defaultdict(int)
                for row_number, row in enumerate(iterator, start=2):
                    values = {key: row[index] for key, index in columns.items()}
                    if all(value is None for value in values.values()):
                        continue
                    label, name, text = values["Category"], values["Advertiser"], values["Ad_copy"]
                    if not all(isinstance(value, str) and value.strip() for value in (label, name, text)):
                        excluded.append({"dataset": dataset, "row": row_number, "reason": "missing_text_name_or_category"})
                        continue
                    category = _ATVI_ALIASES.get(identity(label), known.get(identity(label)))
                    if category not in configured:
                        omitted[label] += 1
                        continue
                    rows[category].append({"ad_id": row_number, "ad_name": name, "ad_copy": text,
                                           "source_record_id": f"dataset/ATVI/ad_data.xlsx#{sheet.title}:{row_number}"})
                    labels[category].add(label)
                info = _file_info(workbook, code)
                for category in rows:
                    sources[category] = [info]
                excluded.extend({"dataset": dataset, "raw_category": label, "records": count,
                                 "reason": "no_existing_query_configuration"} for label, count in sorted(omitted.items()))
            finally:
                book.close()
            mode = "original_ATVI_workbook"
        else:
            mode = "provided_category_JSON_records"
            for category in configured:
                for path in sorted((code / "dataset" / dataset / category).glob("*_section*.json")):
                    sources[category].append(_file_info(path, code))
                    for index, original in enumerate(json.loads(path.read_text(encoding="utf-8"))):
                        record = deepcopy(original)
                        record["source_record_id"] = f"{path.relative_to(code).as_posix()}#{index}"
                        rows[category].append(record)
        for category in configured:
            unique, seen, invalid, duplicates = [], set(), 0, 0
            for record in rows[category]:
                name = record.get("ad_name")
                text = record.get("content", record.get("ad_copy"))
                if not all(isinstance(value, str) and value.strip() for value in (name, text)):
                    invalid += 1
                    continue
                key = identity(text)
                if key in seen:
                    duplicates += 1
                    continue
                seen.add(key)
                unique.append(record)
            pools.append({"dataset": dataset, "category": category, "records": unique,
                          "source": {"group": f"{dataset}/{category}", "files": sources[category],
                                     "source_mode": mode, "raw_category_labels": sorted(labels[category]),
                                     "input_records": len(rows[category]), "duplicate_creatives_removed": duplicates,
                                     "invalid_records_removed": invalid,
                                     "distinct_creatives": len(unique),
                                     "distinct_advertisers": len({identity(r["ad_name"]) for r in unique})}})
    return pools, excluded


def sample_pool(pool, rng, *, n=20):
    if isinstance(n, bool) or not isinstance(n, (int, np.integer)) or n < 2:
        raise ValueError("n must be an integer of at least two distinct advertisers")
    records = pool["records"]
    if pool["source"]["distinct_advertisers"] < n:
        raise ValueError(f'{pool["source"]["group"]}: {pool["source"]["distinct_advertisers"]} distinct advertisers; need {n}')
    unit = "characters" if pool["dataset"] == "ADGEN" else "words"
    lengths = np.array([count_text(r.get("content", r.get("ad_copy")), unit) for r in records])
    if np.any(lengths <= 0):
        raise ValueError("Candidate pool contains an empty creative")
    edges = np.histogram_bin_edges(lengths, bins=INTERVALS)
    bins = np.clip(np.searchsorted(edges, lengths, side="right") - 1, 0, INTERVALS - 1)
    by_bin = defaultdict(lambda: defaultdict(list))
    for index, (record, bin_id) in enumerate(zip(records, bins)):
        by_bin[int(bin_id)][identity(record["ad_name"])].append(index)
    order = rng.permutation(sorted(by_bin)).tolist()
    choices = {bin_id: rng.permutation(sorted(by_bin[bin_id])).tolist() for bin_id in order}
    slots, advertiser_owner = [], {}

    def augment(slot, visited):
        for advertiser in choices[slots[slot]]:
            if advertiser in visited:
                continue
            visited.add(advertiser)
            previous = advertiser_owner.get(advertiser)
            if previous is None or augment(previous, visited):
                advertiser_owner[advertiser] = slot
                return True
        return False

    while len(slots) < n:
        before = len(slots)
        for bin_id in order:
            slots.append(bin_id)
            if not augment(len(slots) - 1, set()):
                slots.pop()
            if len(slots) == n:
                break
        if len(slots) == before:
            raise ValueError("Cannot fill the requested unique-advertiser pool without replacement")
    owners = {slot: advertiser for advertiser, slot in advertiser_owner.items()}
    selected = []
    for slot, bin_id in enumerate(slots):
        index = int(rng.choice(by_bin[bin_id][owners[slot]]))
        record = deepcopy(records[index])
        record.update(length_interval=bin_id, source_length=int(lengths[index]))
        selected.append(record)
    source = deepcopy(pool["source"])
    source.update(pool_policy=POLICY, requested_n=int(n), interval_count=INTERVALS,
                  length_unit=unit, bin_edges=edges.tolist(),
                  bin_boundary_policy="left_closed_right_open_except_last_right_closed",
                  population_bin_counts=np.bincount(bins, minlength=INTERVALS).tolist(),
                  sampled_bin_counts=np.bincount(slots, minlength=INTERVALS).tolist(),
                  sampling_rule="randomized_balanced_bin_passes_unique_advertiser_matching",
                  identity_rule="NFKC_whitespace_casefold_for_names_and_creative_deduplication")
    return selected, source
